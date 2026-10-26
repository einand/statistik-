"""Download and normalize Valmyndighetens final 2022 and 2026 election workbooks."""
from pathlib import Path
from urllib.request import Request, urlopen
import json
import openpyxl

ROOT = Path(__file__).resolve().parent
FILES = {
    2022: {
        "Riksdag": "https://www.val.se/download/18.162047b519a91d0533118f4b/1764336897948/Roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-riksdagsvalet-2022.xlsx",
        "Region": "https://www.val.se/download/18.162047b519a91d0533118f4d/1764336772618/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-regionval-2022.xlsx",
        "Kommun": "https://www.val.se/download/18.162047b519a91d0533118f4e/1764337121617/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-kommunval-2022.xlsx",
    },
    2026: {
        "Riksdag": "https://www.val.se/download/18.7faaad3f1a0b0c300e458a/1790069872234/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-riksdagsvalet-2026.xlsx",
        "Region": "https://www.val.se/download/18.7faaad3f1a0b0c300e4183a/1791371611110/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-regionvalen-2026.xlsx",
        "Kommun": "https://www.val.se/download/18.7faaad3f1a0b0c300e41823/1791364806507/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-kommunvalen-2026-.xlsx",
    },
}

def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = next(it)
    return header, it

elections = {}
for year, year_files in FILES.items():
 for kind, url in year_files.items():
    path = ROOT / f"val{year}-{kind.lower()}.xlsx"
    if not path.exists():
        req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
        path.write_bytes(urlopen(req, timeout=90).read())
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if year == 2022:
        sheet={"Riksdag":"roster_RD","Region":"roster_RF","Kommun":"roster_KF"}[kind]
        header,it=rows(wb[sheet]); col={str(v).strip():i for i,v in enumerate(header)}
        party_votes={}; valid_votes={}; voter_totals={}; seen_districts=set(); party_names=set(); cast_by_muni={}
        ignored={"Summa giltiga röster","Valdeltagande","ej anmält deltagande","blanka röster","övriga ogiltiga"}
        for row in it:
            district=str(row[col["Distrikt"]]).strip();parts=district.split("-")
            if len(parts)<4: continue
            municipality=parts[1].zfill(2)+parts[2].zfill(2)
            category=str(row[col["Parti"]]).strip();votes=int(row[col["Röster"]] or 0);eligible=int(row[col["Röstberättigade"]] or 0)
            if eligible and district not in seen_districts:
                voter_totals.setdefault(municipality,{"votes":0,"eligible":0})["eligible"]+=eligible;seen_districts.add(district)
            if category=="Summa giltiga röster": valid_votes[municipality]=valid_votes.get(municipality,0)+votes
            elif category=="Valdeltagande": cast_by_muni[municipality]=cast_by_muni.get(municipality,0)+votes
            elif category not in ignored and "ogiltig" not in category.lower():
                party_names.add(category); party_votes.setdefault(municipality,{})[category]=party_votes.setdefault(municipality,{}).get(category,0)+votes
        shares={}
        for municipality,values in party_votes.items():
            denominator=valid_votes.get(municipality,0)
            if denominator: shares[municipality]={party:votes/denominator for party,votes in values.items()}
        turnout={}
        for municipality,totals in voter_totals.items():
            totals["votes"]=cast_by_muni.get(municipality,0)
            if totals["eligible"]: turnout[municipality]=100*totals["votes"]/totals["eligible"]
        total_valid=sum(valid_votes.values()); total_cast=sum(cast_by_muni.values()); total_eligible=sum(v["eligible"] for v in voter_totals.values())
        national={party:sum(party_votes.get(m,{}).get(party,0) for m in party_votes)/total_valid for party in party_names} if total_valid else {}
        shares["00"]=national;turnout["00"]=100*total_cast/total_eligible if total_eligible else None
        election_year={"parties":sorted(party_names),"shares":shares,"turnout":turnout,"validVotes":valid_votes,"voterTotals":voter_totals,"sourceUrl":url,"municipalityCount":len(party_votes)}
        elections.setdefault(kind,{"name":{"Riksdag":"Riksdagsvalet","Region":"Regionvalen","Kommun":"Kommunvalen"}[kind],"years":{}})["years"][str(year)]=election_year
        wb.close();path.unlink();continue
    share_header, share_rows = rows(wb["Andel_kommun"])
    turnout_header, turnout_rows = rows(wb["Valdeltagande_kommun"])
    # The party columns begin after the election/geography identifiers.
    share_party_start = next(i for i, x in enumerate(share_header) if x in ("Arbetarepartiet-Socialdemokraterna", "Socialdemokraterna"))
    parties = [str(x) for x in share_header[share_party_start:] if x]
    municipality = {}
    for row in share_rows:
        code = str(row[share_header.index("Kommunkod")]).zfill(4)
        municipality[code] = {parties[i]: float(row[share_party_start+i] or 0) for i in range(len(parties))}
    turnout_code_i = turnout_header.index("Kommunkod")
    turnout_rate_i = turnout_header.index("Valdeltagande (%)")
    turnout = {str(row[turnout_code_i]).zfill(4): float(row[turnout_rate_i]) * 100
               for row in turnout_rows if row[turnout_rate_i] is not None}
    # Aggregate the national party shares and turnout from municipality totals
    # weighted by valid votes / eligible voters, respectively.
    count_header, count_rows = rows(wb["Antal_kommun"])
    count_party_start = next(i for i, x in enumerate(count_header) if x in parties)
    total_by_party = {p: 0 for p in parties}
    valid_votes = {}
    for row in count_rows:
        code_i = count_header.index("Kommunkod")
        code = str(row[code_i]).zfill(4)
        valid_votes[code] = sum(float(row[count_party_start+i] or 0) for i in range(len(parties)) if count_party_start+i < len(row))
        for i, party in enumerate(parties):
            if count_party_start+i < len(row) and row[count_party_start+i] is not None:
                total_by_party[party] += float(row[count_party_start+i])
    party_sum = sum(total_by_party.values())
    national = {p: v / party_sum for p, v in total_by_party.items()} if party_sum else {}
    rate_col = turnout_header.index("Valdeltagande (%)")
    votes_col = turnout_header.index("Röster")
    eligible_col = turnout_header.index("Röstberättigade")
    votes = eligible = 0
    voter_totals = {}
    for row in wb["Valdeltagande_kommun"].iter_rows(min_row=2, values_only=True):
        code = str(row[turnout_code_i]).zfill(4)
        local_votes, local_eligible = float(row[votes_col] or 0), float(row[eligible_col] or 0)
        voter_totals[code] = {"votes":local_votes,"eligible":local_eligible}
        votes += local_votes
        eligible += local_eligible
    turnout["00"] = 100 * votes / eligible if eligible else None
    # County values from Valmyndighetens own county-level results tabs.
    county_share = {}
    county_turnout = {}
    if "Andel_län" in wb.sheetnames:
        h, it = rows(wb["Andel_län"])
        start = next(i for i, x in enumerate(h) if x in parties)
        code_i = next(i for i, x in enumerate(h) if x in ("Länskod", "Regionkod"))
        for row in it:
            code = str(row[code_i]).zfill(2)
            if code == "99": continue
            county_share[code] = {p: float(row[start+i] or 0) for i, p in enumerate(parties)}
    if "Valdeltagande_län" in wb.sheetnames:
        h, it = rows(wb["Valdeltagande_län"])
        code_i = next(i for i, x in enumerate(h) if x in ("Länskod", "Regionkod"))
        rate_i = h.index("Valdeltagande (%)")
        for row in it:
            code = str(row[code_i]).zfill(2)
            if code != "99" and row[rate_i] is not None:
                county_turnout[code] = float(row[rate_i]) * 100
    election_year = {
        "parties": parties,
        "shares": {"00": national, **county_share, **municipality},
        "turnout": {**county_turnout, **turnout},
        "validVotes": valid_votes,
        "voterTotals": voter_totals,
        "sourceUrl": url,
        "municipalityCount": len(municipality),
    }
    elections.setdefault(kind,{"name": {"Riksdag":"Riksdagsvalet", "Region":"Regionvalen", "Kommun":"Kommunvalen"}[kind],"years":{}})["years"][str(year)] = election_year
    wb.close()
    path.unlink()

(ROOT / "val-election.json").write_text(json.dumps(elections, ensure_ascii=False, separators=(",", ":")))
print("Valmyndigheten:", "; ".join(f"{year} {kind}: {data['municipalityCount']} kommuner, {len(data['parties'])} partier" for kind,e in elections.items() for year,data in e['years'].items()))
